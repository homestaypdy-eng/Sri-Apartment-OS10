"""Focused Expense payment-ledger, cash, P&L, RBAC, and audit regressions."""
from concurrent.futures import ThreadPoolExecutor
from collections import Counter
from datetime import date, timedelta
from io import BytesIO
import uuid
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.modules.finance import router as finance_router
from conftest import ADMIN_EMAIL, ADMIN_PASSWORD

client = TestClient(app)


@pytest.fixture(scope="module")
def admin_headers():
    return _headers()


def _headers(email=ADMIN_EMAIL, password=ADMIN_PASSWORD):
    response = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _new_user(admin_headers, role):
    identity = uuid.uuid4().hex[:12]
    email = f"expense_{role}_{identity}@sriapartmenthomestay.com"
    created = client.post("/api/v1/auth/register", headers=admin_headers, json={
        "name": f"Expense {role.title()} {identity}", "email": email,
        "password": "ExpenseTest123!", "role": role,
    })
    assert created.status_code == 201, created.text
    return _headers(email, "ExpenseTest123!")


def _create_expense(headers, *, amount="100.00", status="approved", expense_date=None, category=None, unit_id=None):
    response = client.post("/api/v1/finance/expenses", headers=headers, json={
        "category": category or f"Ledger-{uuid.uuid4().hex[:8]}",
        "amount": amount,
        "expense_date": (expense_date or date.today()).isoformat(),
        "approval_status": status,
        "unit_id": unit_id,
    })
    assert response.status_code == 201, response.text
    return response.json()


def _pay(headers, expense_id, amount, payment_date=None, method="cash", reference=None, idempotency_key=None):
    request_headers = {**headers, "Idempotency-Key": idempotency_key or uuid.uuid4().hex}
    return client.post(f"/api/v1/finance/expenses/{expense_id}/payments", headers=request_headers, json={
        "amount": amount,
        "payment_date": (payment_date or date.today()).isoformat(),
        "payment_method": method,
        "reference_no": reference,
    })


def _unit(headers):
    prop = client.post("/api/v1/properties", headers=headers, json={
        "name": f"Expense Ledger Property {uuid.uuid4().hex[:6]}", "address": "Test", "state_code": "34",
    })
    assert prop.status_code == 201, prop.text
    unit = client.post("/api/v1/units", headers=headers, json={
        "property_id": prop.json()["id"], "name": f"Expense Ledger Unit {uuid.uuid4().hex[:6]}", "base_rate": 1500,
    })
    assert unit.status_code == 201, unit.text
    return prop.json()["id"], unit.json()["id"]


def test_full_and_partial_payments_derive_balance_and_are_audited_once(admin_headers):
    admin = admin_headers
    manager = _new_user(admin, "manager")
    expense = _create_expense(admin, amount="125.50")

    first = _pay(manager, expense["id"], "25.50", method="upi", reference="UT-1")
    assert first.status_code == 201, first.text
    listed = client.get("/api/v1/finance/expenses", headers=admin).json()
    row = next(item for item in listed if item["id"] == expense["id"])
    assert (row["paid_amount"], row["remaining_amount"], row["payment_status"]) == (25.5, 100.0, "partially_paid")

    second = _pay(admin, expense["id"], "100.00", method="cheque", reference="CH-2")
    assert second.status_code == 201, second.text
    row = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == expense["id"])
    assert (row["paid_amount"], row["remaining_amount"], row["payment_status"]) == (125.5, 0.0, "paid")
    assert row["approval_status"] == "approved"
    assert row["amount"] == 125.5
    assert client.patch(f"/api/v1/finance/expenses/{expense['id']}", headers=admin, json={
        "category": "Edited", "amount": "100", "expense_date": date.today().isoformat(),
    }).status_code == 409
    assert client.delete(f"/api/v1/finance/expenses/{expense['id']}", headers=admin).status_code == 409

    history = client.get(f"/api/v1/finance/expenses/{expense['id']}/payments", headers=admin)
    assert history.status_code == 200
    assert [item["payment_method"] for item in history.json()] == ["upi", "cheque"]
    assert history.json()[0]["reference_no"] == "UT-1"
    assert history.json()[0]["created_by"] is not None
    logs = client.get("/api/v1/audit/logs", headers=admin, params={
        "action": "finance.expense.payment.create", "entity_type": "expense_payment",
    })
    assert logs.status_code == 200
    matching = [item for item in logs.json() if item["entity_id"] in {first.json()["id"], second.json()["id"]}]
    assert len(matching) == 2
    assert Counter(item["entity_id"] for item in matching) == Counter({first.json()["id"]: 1, second.json()["id"]: 1})
    assert all(item["user_id"] is not None and item["new_value"] for item in matching)


def test_payment_validation_rbac_and_approval_separation(monkeypatch, admin_headers):
    admin = admin_headers
    manager = _new_user(admin, "manager")
    staff = _new_user(admin, "staff")
    expense = _create_expense(admin, amount="40.00", status="pending")
    path = f"/api/v1/finance/expenses/{expense['id']}/payments"

    assert _pay(staff, expense["id"], "10").status_code == 403
    assert _pay({}, expense["id"], "10").status_code == 401
    assert _pay(manager, expense["id"], "0").status_code == 422
    assert _pay(manager, expense["id"], "-1").status_code == 422
    assert _pay(manager, expense["id"], "10", method="wire").status_code == 422
    assert _pay(manager, expense["id"], "41").status_code == 409
    assert _pay(manager, str(uuid.uuid4()), "1").status_code == 404
    assert client.get(path, headers=staff).status_code == 200

    assert _pay(manager, expense["id"], "10").status_code == 409
    approved = client.post(f"/api/v1/finance/expenses/{expense['id']}/approval", headers=admin, json={"approve": True})
    assert approved.status_code == 200, approved.text

    success = _pay(manager, expense["id"], "10", method="bank_transfer")
    assert success.status_code == 201
    spoofed = client.post(
        path,
        headers={**manager, "Idempotency-Key": uuid.uuid4().hex},
        json={"amount": "1", "payment_date": date.today().isoformat(), "payment_method": "cash", "created_by": str(uuid.uuid4())},
    )
    assert spoofed.status_code == 201
    assert spoofed.json()["created_by"] == success.json()["created_by"]
    current = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == expense["id"])
    assert current["approval_status"] == "approved"
    assert current["amount"] == 40.0
    # 10 (success) + 1 (spoofed created_by, otherwise a legitimate payment
    # that must still succeed) — the spoof attempt only ignores created_by,
    # it doesn't block the payment itself.
    assert current["paid_amount"] == 11.0

    rejected = _create_expense(admin, amount="10.00", status="pending")
    monkeypatch.setattr(finance_router, "trigger_expense_approval", lambda *args, **kwargs: None)
    decision = client.post(f"/api/v1/finance/expenses/{rejected['id']}/approval", headers=admin, json={"approve": False})
    assert decision.status_code == 200
    recorded = _pay(manager, rejected["id"], "5", method="card")
    assert recorded.status_code == 409
    current_rejected = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == rejected["id"])
    assert current_rejected["approval_status"] == "rejected"


def test_cash_flow_cashbook_and_new_daily_closing_use_payment_date_only(monkeypatch, admin_headers):
    admin = admin_headers
    payment_day = date(2500 + uuid.uuid4().int % 1000, 7, 16)
    expense = _create_expense(admin, amount="60.00", expense_date=date(2040, 7, 1))
    cash_before = client.get("/api/v1/finance/cash-flow", headers=admin, params={"date_from": payment_day.isoformat(), "date_to": payment_day.isoformat()}).json()

    monkeypatch.setattr(finance_router, "trigger_daily_closing_completed", lambda *args, **kwargs: None)
    closed = client.post("/api/v1/finance/daily-closing", headers=admin, json={"closing_date": payment_day.isoformat(), "opening_balance": 100})
    assert closed.status_code == 201, closed.text
    snapshot = closed.json()

    payment = _pay(admin, expense["id"], "35.25", payment_day, method="bank_transfer")
    assert payment.status_code == 201, payment.text
    cash_after = client.get("/api/v1/finance/cash-flow", headers=admin, params={"date_from": payment_day.isoformat(), "date_to": payment_day.isoformat()}).json()
    assert round(cash_after["cash_out_expenses"] - cash_before["cash_out_expenses"], 2) == 35.25
    cashbook = client.get("/api/v1/finance/cashbook", headers=admin, params={"cashbook_date": payment_day.isoformat()}).json()
    line = next(line for line in cashbook["lines"] if "Expense payment" in line["description"])
    assert line["amount"] == 35.25 and "bank_transfer" in line["description"]

    stored = client.get("/api/v1/finance/daily-closing", headers=admin, params={"date_from": payment_day.isoformat(), "date_to": payment_day.isoformat()}).json()
    assert len(stored) == 1
    assert stored[0]["total_cash_out"] == snapshot["total_cash_out"]
    assert stored[0]["closing_balance"] == snapshot["closing_balance"]

    next_day = payment_day + timedelta(days=1)
    second_payment = _pay(admin, expense["id"], "10.00", next_day, method="cash")
    assert second_payment.status_code == 201, second_payment.text
    next_cashbook = client.get("/api/v1/finance/cashbook", headers=admin, params={"cashbook_date": next_day.isoformat()}).json()
    next_close = client.post("/api/v1/finance/daily-closing", headers=admin, json={"closing_date": next_day.isoformat(), "opening_balance": 200})
    assert next_close.status_code == 201, next_close.text
    assert next_close.json()["total_cash_out"] == next_cashbook["total_out"]


def test_pnl_dashboard_maintenance_and_finance_bundle_exclude_rejected(monkeypatch, admin_headers):
    admin = admin_headers
    property_id, unit_id = _unit(admin)
    today = date.today()
    category = f"Policy-{uuid.uuid4().hex[:8]}"
    before_dashboard = client.get("/api/v1/reports/dashboard", headers=admin).json()
    before_summary = client.get("/api/v1/reports/summary", headers=admin, params={
        "date_from": today.replace(day=1).isoformat(), "date_to": (today + timedelta(days=1)).isoformat(),
    }).json()
    before_apartment = client.get(f"/api/v1/finance/apartment-pnl/{unit_id}", headers=admin, params={"period_year": today.year, "period_month": today.month}).json()
    before_property = client.get(f"/api/v1/finance/property-pnl/{property_id}", headers=admin, params={"period_year": today.year, "period_month": today.month}).json()
    before_bundle = client.get("/api/v1/reports/ai-strategist", headers=admin, params={
        "date_from": today.replace(day=1).isoformat(), "date_to": (today + timedelta(days=1)).isoformat(), "property_id": property_id,
    })
    assert before_bundle.status_code == 200, before_bundle.text

    pending_unit = _create_expense(admin, amount="11", status="pending", expense_date=today, category=category, unit_id=unit_id)
    approved_unit = _create_expense(admin, amount="12", status="approved", expense_date=today, category=category, unit_id=unit_id)
    rejected_unit = _create_expense(admin, amount="13", status="pending", expense_date=today, category=category, unit_id=unit_id)
    central_pending = _create_expense(admin, amount="21", status="pending", expense_date=today, category=category)
    central_approved = _create_expense(admin, amount="22", status="approved", expense_date=today, category=category)
    central_rejected = _create_expense(admin, amount="23", status="pending", expense_date=today, category=category)
    maintenance_pending = _create_expense(admin, amount="31", status="pending", expense_date=today, category="maintenance", unit_id=unit_id)
    maintenance_approved = _create_expense(admin, amount="32", status="approved", expense_date=today, category="maintenance", unit_id=unit_id)
    maintenance_rejected = _create_expense(admin, amount="33", status="pending", expense_date=today, category="maintenance", unit_id=unit_id)

    monkeypatch.setattr(finance_router, "trigger_expense_approval", lambda *args, **kwargs: None)
    for expense in (rejected_unit, central_rejected, maintenance_rejected):
        result = client.post(f"/api/v1/finance/expenses/{expense['id']}/approval", headers=admin, json={"approve": False})
        assert result.status_code == 200

    budget = client.get("/api/v1/finance/budget-vs-actual", headers=admin, params={"period_year": today.year, "period_month": today.month}).json()
    category_line = next(row for row in budget["lines"] if row["category"] == category)
    assert category_line["actual_amount"] == 66.0
    apartment = client.get(f"/api/v1/finance/apartment-pnl/{unit_id}", headers=admin, params={"period_year": today.year, "period_month": today.month}).json()
    assert apartment["total_expenses"] - before_apartment["total_expenses"] == 86.0  # 11 + 12 + 31 + 32; rejected rows excluded
    prop_pnl = client.get(f"/api/v1/finance/property-pnl/{property_id}", headers=admin, params={"period_year": today.year, "period_month": today.month}).json()
    assert prop_pnl["central_expenses"] - before_property["central_expenses"] == 43.0
    assert prop_pnl["total_expenses"] - before_property["total_expenses"] == 129.0

    after_dashboard = client.get("/api/v1/reports/dashboard", headers=admin).json()
    assert after_dashboard["today_expenses"] - before_dashboard["today_expenses"] == 129.0
    assert after_dashboard["month_expenses"] - before_dashboard["month_expenses"] == 129.0
    unit_dashboard = next(row for row in after_dashboard["apartment_performance"] if row["unit_id"] == unit_id)
    before_unit = next((row for row in before_dashboard["apartment_performance"] if row["unit_id"] == unit_id), None)
    assert unit_dashboard["month_expenses"] - (before_unit["month_expenses"] if before_unit else 0) == 86.0
    trend = next(row for row in after_dashboard["monthly_trend"] if row["year"] == today.year and row["month"] == today.month)
    before_trend = next(row for row in before_dashboard["monthly_trend"] if row["year"] == today.year and row["month"] == today.month)
    assert trend["expenses"] - before_trend["expenses"] == 129.0

    after_summary = client.get("/api/v1/reports/summary", headers=admin, params={
        "date_from": today.replace(day=1).isoformat(), "date_to": (today + timedelta(days=1)).isoformat(),
    }).json()
    assert after_summary["maintenance_cost_total"] - before_summary["maintenance_cost_total"] == 63.0
    assert after_summary["revenue"]["total"] == before_summary["revenue"]["total"]

    bundle = client.get("/api/v1/reports/ai-strategist", headers=admin, params={
        "date_from": today.replace(day=1).isoformat(), "date_to": (today + timedelta(days=1)).isoformat(),
        "property_id": property_id,
    })
    assert bundle.status_code == 200, bundle.text
    finance = bundle.json()["finance"]
    before_finance = before_bundle.json()["finance"]
    assert finance["pnl"]["central_expenses"] - before_finance["pnl"]["central_expenses"] == 43.0
    assert finance["pnl"]["operating_expenses"] - before_finance["pnl"]["operating_expenses"] == 129.0

    listed = client.get("/api/v1/finance/expenses", headers=admin, params={"date_from": today.isoformat(), "date_to": today.isoformat(), "category": category}).json()
    ids = {row["id"] for row in listed}
    assert {pending_unit["id"], approved_unit["id"], rejected_unit["id"], central_pending["id"], central_approved["id"], central_rejected["id"]} <= ids
    exported = client.get("/api/v1/finance/reports/export", headers=admin, params={
        "report": "expenses", "format": "xlsx", "date_from": today.isoformat(), "date_to": today.isoformat(),
    })
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
    exported_statuses = {str(row[4]) for row in workbook.active.iter_rows(min_row=2, values_only=True) if row[1] == category}
    assert {"pending", "approved", "rejected"} <= exported_statuses
    workbook.close()


def test_concurrent_partial_payments_cannot_overpay(admin_headers):
    admin = admin_headers
    expense = _create_expense(admin, amount="100.00")
    barrier = Barrier(2)

    def attempt():
        with TestClient(app) as concurrent_client:
            barrier.wait(timeout=5)
            return concurrent_client.post(
                f"/api/v1/finance/expenses/{expense['id']}/payments",
                headers={**admin, "Idempotency-Key": uuid.uuid4().hex},
                json={"amount": "75.00", "payment_date": date.today().isoformat(), "payment_method": "cash"},
            ).status_code

    with ThreadPoolExecutor(max_workers=2) as executor:
        outcomes = list(executor.map(lambda _: attempt(), range(2)))
    assert sorted(outcomes) == [201, 409]
    final = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == expense["id"])
    assert final["paid_amount"] == 75.0


def test_idempotency_retry_returns_original_payment_and_writes_one_audit(admin_headers):
    admin = admin_headers
    expense = _create_expense(admin, amount="100")
    key = uuid.uuid4().hex
    first = _pay(admin, expense["id"], "25.00", method="upi", idempotency_key=key)
    replay = _pay(admin, expense["id"], "25.00", method="upi", idempotency_key=key)
    assert first.status_code == 201, first.text
    assert replay.status_code == 200, replay.text
    assert replay.json()["id"] == first.json()["id"]
    history = client.get(f"/api/v1/finance/expenses/{expense['id']}/payments", headers=admin).json()
    assert len([p for p in history if p["id"] == first.json()["id"]]) == 1
    audit = client.get("/api/v1/audit/logs", headers=admin, params={"entity_type": "expense_payment", "entity_id": first.json()["id"]})
    assert audit.status_code == 200, audit.text
    assert len([item for item in audit.json() if item["action"] == "finance.expense.payment.create"]) == 1


def test_idempotency_key_rejects_payload_or_expense_reuse_but_new_key_allows_second_payment(admin_headers):
    admin = admin_headers
    first_expense = _create_expense(admin, amount="100")
    second_expense = _create_expense(admin, amount="100")
    key = uuid.uuid4().hex
    first = _pay(admin, first_expense["id"], "20", idempotency_key=key)
    assert first.status_code == 201, first.text
    changed_amount = _pay(admin, first_expense["id"], "21", idempotency_key=key)
    changed_expense = _pay(admin, second_expense["id"], "20", idempotency_key=key)
    assert changed_amount.status_code == 409
    assert changed_expense.status_code == 409
    second_payment = _pay(admin, first_expense["id"], "30", idempotency_key=uuid.uuid4().hex)
    assert second_payment.status_code == 201, second_payment.text
    listed = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == first_expense["id"])
    assert listed["paid_amount"] == 50


def test_legacy_payment_row_with_no_idempotency_key_stays_compatible(admin_headers):
    """A payment row created before idempotency_key/request_fingerprint
    existed (both columns are nullable, additive) must still: sum
    correctly into paid_amount/remaining_amount, appear in the payment
    history list, and coexist with a normal new API-created payment for
    the same expense — the new (non-partial) unique index on
    idempotency_key must not treat two NULLs as a collision (Postgres
    itself never does, but this proves the whole stack agrees)."""
    admin = admin_headers
    expense = _create_expense(admin, amount="100.00")

    from app.db.session import SessionLocal
    from app.modules.finance.models import ExpensePayment
    db = SessionLocal()
    try:
        legacy = ExpensePayment(
            expense_id=uuid.UUID(expense["id"]), amount=30, payment_date=date.today(),
            payment_method="cash", reference_no="LEGACY-PRE-IDEMPOTENCY",
            # idempotency_key / request_fingerprint deliberately left unset (NULL) —
            # exactly what a row from before this migration looks like.
        )
        db.add(legacy)
        db.commit()
        legacy_id = str(legacy.id)
    finally:
        db.close()

    row = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == expense["id"])
    assert (row["paid_amount"], row["remaining_amount"], row["payment_status"]) == (30.0, 70.0, "partially_paid")
    history = client.get(f"/api/v1/finance/expenses/{expense['id']}/payments", headers=admin).json()
    assert any(item["id"] == legacy_id for item in history)

    # A brand-new, idempotency-keyed payment must still work fine alongside
    # the legacy NULL-key row.
    new_payment = _pay(admin, expense["id"], "20.00", method="upi")
    assert new_payment.status_code == 201, new_payment.text
    row = next(item for item in client.get("/api/v1/finance/expenses", headers=admin).json() if item["id"] == expense["id"])
    assert row["paid_amount"] == 50.0


def test_concurrent_idempotent_retry_creates_one_payment(admin_headers):
    admin = admin_headers
    expense = _create_expense(admin, amount="100")
    key = uuid.uuid4().hex
    barrier = Barrier(2)

    def attempt():
        with TestClient(app) as concurrent_client:
            barrier.wait(timeout=5)
            return concurrent_client.post(
                f"/api/v1/finance/expenses/{expense['id']}/payments",
                headers={**admin, "Idempotency-Key": key},
                json={"amount": "25.00", "payment_date": date.today().isoformat(), "payment_method": "cash", "reference_no": None},
            )

    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(executor.map(lambda _: attempt(), range(2)))
    assert sorted(response.status_code for response in responses) == [200, 201]
    assert responses[0].json()["id"] == responses[1].json()["id"]
    history = client.get(f"/api/v1/finance/expenses/{expense['id']}/payments", headers=admin).json()
    assert len(history) == 1
